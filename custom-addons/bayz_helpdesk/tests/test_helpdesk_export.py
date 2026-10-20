import base64
from datetime import date, datetime
from io import BytesIO

from openpyxl import load_workbook

from odoo.exceptions import ValidationError
from odoo.tests.common import TransactionCase

from .common import HelpdeskTestMixin


# Validation des filtres et contenu du fichier Excel exporté.
class TestHelpdeskTimesheetExport(HelpdeskTestMixin, TransactionCase):

    def setUp(self):
        super().setUp()
        self.technician = self.create_user(
            "Export technician", ["bayz_helpdesk.group_helpdesk_technician"]
        )
        self.category = self.env["helpdesk.category"].create({
            "name": "Export category",
        })

    def test_date_range_must_be_valid(self):
        with self.assertRaises(ValidationError):
            self.env["helpdesk.timesheet.export.wizard"].create({
                "date_from": date(2025, 2, 2),
                "date_to": date(2025, 2, 1),
            })

    def test_export_contains_filtered_technician_statistics(self):
        ticket = self.create_ticket(category_id=self.category.id)
        self.env["helpdesk.timesheet"].create({
            "ticket_id": ticket.id,
            "technician_id": self.technician.id,
            "start_datetime": datetime(2025, 1, 15, 12, 0),
            "end_datetime": datetime(2025, 1, 15, 14, 30),
        })
        wizard = self.env["helpdesk.timesheet.export.wizard"].create({
            "date_from": date(2025, 1, 15),
            "date_to": date(2025, 1, 15),
            "technician_ids": [(6, 0, [self.technician.id])],
            "category_ids": [(6, 0, [self.category.id])],
        })

        action = wizard.action_export_excel()

        workbook = load_workbook(BytesIO(base64.b64decode(wizard.file_data)))
        sheet = workbook["Statistiques"]
        rows = list(sheet.iter_rows(values_only=True))

        self.assertEqual(action["type"], "ir.actions.act_url")
        self.assertTrue(wizard.file_name.endswith("20250115_20250115.xlsx"))
        self.assertEqual(rows[0], (
            "Technicien",
            "Nombre de tickets",
            "Temps total (h)",
            "Temps moyen / ticket (h)",
        ))
        self.assertEqual(rows[1], (self.technician.name, 1, 2.5, 2.5))
